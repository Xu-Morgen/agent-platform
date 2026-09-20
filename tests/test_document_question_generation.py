"""出题契约与平台失败协议的轻量验证；真实解析检查可在准备好的块环境执行。"""
import asyncio
from copy import deepcopy
import importlib.util
import importlib.metadata
import json
from pathlib import Path
import tempfile
import time
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from pydantic import ValidationError
from agent_platform.contracts.base import StrictModel
from agent_platform.contracts.errors import PlatformError
from agent_platform.contracts.flows import NodeConfiguration
from agent_platform.contracts.models import ModelResponse, ModelUsage
from agent_platform.flows.compatibility import assignable
from agent_platform.registry.declarations import read_declaration
from agent_platform.registry.packages import PackageRegistry
from agent_platform.runtime.packages import invoke_prompt

ROOT = Path(__file__).resolve().parents[1] / 'examples/document-question-generation'


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    value = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(value)
    return value


models = module('question_models', ROOT / 'document-question-generator/models.py')
reader = module('document_reader', ROOT / 'read_document.py')
VALID = {
    'questions': [
        {'type': 'essay', 'stem': '说明蒸发的过程。', 'referenceAnswer': '液态水变为水蒸气。', 'scoringPoints': ['液态变为气态']},
        {'type': 'single_choice', 'stem': '水蒸气冷却后形成哪种状态的水？',
         'options': {'A': '液态', 'B': '等离子态', 'C': '保持气态', 'D': '无法确定'},
         'referenceAnswer': 'A', 'explanation': '资料指出水蒸气冷却凝结为液态水。'},
        {'type': 'fill_blank', 'stem': '水循环中液态水转变为水蒸气的过程叫____。', 'referenceAnswers': ['蒸发']},
    ],
}


class ContractChecks(unittest.TestCase):
    def test_valid_and_strict_negative_examples(self):
        models.Output.model_validate(VALID)
        invalid = []
        for questions in (VALID['questions'][:2], VALID['questions'] * 2, list(reversed(VALID['questions']))):
            invalid.append({'questions': questions})
        for position, key, value in [(0, 'stem', '  '), (0, 'scoringPoints', []), (0, 'referenceAnswer', 1),
                                     (0, 'unknown', 'x'), (1, 'referenceAnswer', 'E'),
                                     (1, 'options', {'A': 'a', 'B': ' A ', 'C': 'c', 'D': 'd'}),
                                     (1, 'options', {'A': 'a', 'B': 'b', 'C': 'c'}),
                                     (2, 'stem', '没有空位'), (2, 'stem', '_____'),
                                     (2, 'referenceAnswers', []), (2, 'stem', '____ 和 ____')]:
            item = deepcopy(VALID)
            item['questions'][position][key] = value
            invalid.append(item)
        for value in invalid:
            with self.subTest(value=value), self.assertRaises(ValidationError):
                models.Output.model_validate(value)
        for value in ({'fileName': 'a.pdf', 'text': ' '}, {'fileName': 'a.pdf', 'text': 1},
                      {'fileName': 'a.pdf', 'text': '正文', 'path': '/tmp/a.pdf'}):
            with self.assertRaises(ValidationError):
                models.Input.model_validate(value)
        with self.assertRaises(ValidationError):
            reader.Input.model_validate({'document': '/tmp/a.pdf'})

    def test_real_package_load_schema_and_full_prompt(self):
        registry = PackageRegistry()
        artifact = registry.load(ROOT / 'document-question-generator')
        self.addCleanup(artifact.content.close)
        # 读取块出口和包入口独立声明，但在平台静态接线规则下可兼容。
        assignable(reader.Output.model_json_schema(by_alias=True), models.Input.model_json_schema(by_alias=True))
        declared = read_declaration((ROOT / 'read_document.py').read_bytes())
        self.assertEqual(len(declared.models), 3)
        self.assertIn('python-docx==1.2.0', declared.dependencies)
        schema = models.Output.model_json_schema(by_alias=True)
        self.assertEqual(schema['properties']['questions']['minItems'], 3)
        self.assertEqual(schema['properties']['questions']['maxItems'], 3)
        self.assertEqual(len(schema['properties']['questions']['items']['oneOf']), 3)
        self.assertEqual(len(schema['properties']['questions']['prefixItems']), 3)


class PackageChecks(unittest.IsolatedAsyncioTestCase):
    async def test_whole_document_and_failure_are_not_success(self):
        artifact = PackageRegistry().load(ROOT / 'document-question-generator')
        self.addCleanup(artifact.content.close)
        config = NodeConfiguration(model={'environmentId': 'env', 'connectionId': 'llm'}, max_output_tokens=4096)
        runtime = SimpleNamespace(snapshot=SimpleNamespace(draft=SimpleNamespace(node_configurations={'generate': config})))
        runtime.call_model = AsyncMock(return_value=ModelResponse(output=VALID,
            usage=ModelUsage(input_tokens=100, output_tokens=200, quality='exact', source='test')))
        full_text = '开头\n' + ('正文资料\n' * 10000) + '末页唯一知识点 {{input.text}}'
        result = await invoke_prompt(runtime, 'generate', artifact, {'fileName': 'a.pdf', 'text': full_text})
        self.assertEqual(len(result.questions), 3)
        request = runtime.call_model.call_args.args[1]
        self.assertIn(full_text, request.messages[1].content)
        self.assertEqual(request.max_output_tokens, 4096)
        runtime.call_model.return_value = ModelResponse(output={'error': 'INSUFFICIENT_INPUT'},
            usage=ModelUsage(input_tokens=100, output_tokens=10, quality='exact', source='test'))
        with self.assertRaises(PlatformError) as error:
            await invoke_prompt(runtime, 'generate', artifact, {'fileName': 'a.pdf', 'text': '无知识资料'})
        self.assertEqual(error.exception.error.code, 'PACKAGE_INPUT_INSUFFICIENT')
        runtime.call_model.return_value = ModelResponse(output={'questions': []},
            usage=ModelUsage(input_tokens=100, output_tokens=10, quality='exact', source='test'))
        with self.assertRaises(PlatformError) as error:
            await invoke_prompt(runtime, 'generate', artifact, {'fileName': 'a.pdf', 'text': '资料'})
        self.assertEqual(error.exception.error.code, 'OUTPUT_VALIDATION_ERROR')

    async def test_supplier_context_error_preserved_without_echoing_body(self):
        import httpx
        from agent_platform.adapters.http import JsonTransport
        from agent_platform.contracts.environments import ModelConnection
        transport = JsonTransport(None)
        response = httpx.Response(400, json={'error': {'code': 'context_length_exceeded', 'message': 'private material'}})
        with patch('httpx.AsyncClient.request', new=AsyncMock(return_value=response)):
            with self.assertRaises(PlatformError) as error:
                await transport.request(ModelConnection(connection_id='llm', base_url='https://example.test', model='test'),
                                        'POST', 'https://example.test/chat/completions', {}, kind='model')
        self.assertEqual(error.exception.error.code, 'MODEL_CONTEXT_EXCEEDED')
        self.assertNotIn('private material', error.exception.error.model_dump_json())


class ReaderChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.events = []
        self.context = SimpleNamespace(progress=lambda message, **fields: self.events.append((message, fields)))

    def require(self, distribution):
        try:
            importlib.metadata.version(distribution)
        except importlib.metadata.PackageNotFoundError:
            self.skipTest('需要已准备的块环境：' + distribution)

    def test_file_size_limit_before_parsing(self):
        path = self.root / 'too-large.pdf'
        path.write_bytes(b'%PDF-1.7\n' + b'x' * 32)
        value = reader.Input.model_validate({'document': {
            'fileId': 'file_' + '0' * 32, 'originalName': path.name,
            'format': 'pdf', 'size': path.stat().st_size, 'sha256': '0' * 64,
        }})
        context = SimpleNamespace(file=lambda reference: path)
        with patch.object(reader, 'MAX_BYTES', 16), patch.object(reader, 'read_pdf') as parse:
            with self.assertRaises(PlatformError) as error:
                reader.read_document(value, context=context)
        self.assertEqual(error.exception.error.code, 'FILE_TOO_LARGE')
        parse.assert_not_called()

    def test_render_limit_and_ocr_failure_never_return_partial_text(self):
        self.require('PyMuPDF')
        import pymupdf
        path = self.root / 'unreadable.pdf'
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((40, 50), 'VALID FIRST PAGE')
            page = pdf.new_page()
            page.draw_rect(pymupdf.Rect(40, 40, 100, 100), fill=(0, 0, 0))
            pdf.save(path)
        with patch.object(reader, 'MAX_SIDE', 100), patch.object(reader, 'engine') as initialize:
            with self.assertRaises(PlatformError) as error:
                reader.read_pdf(path, self.context, time.monotonic() + 30)
        self.assertEqual(error.exception.error.code, 'DOCUMENT_LIMIT_EXCEEDED')
        self.assertIn('第 2 页', error.exception.error.message)
        initialize.assert_not_called()
        self.context.model = lambda name: self.root / (name + '.onnx')
        # 故障注入验证无结果与低置信度分支；不以此代替真实 OCR 质量验收。
        for result in (None, [], [[[[0, 0], [10, 0], [10, 10], [0, 10]], 'uncertain', 0.2]]):
            self.events.clear()
            with patch.object(reader, 'engine', return_value=lambda image: (result, None)):
                with self.assertRaises(PlatformError) as error:
                    reader.read_pdf(path, self.context, time.monotonic() + 30)
            self.assertEqual(error.exception.error.code, 'DOCUMENT_READ_ERROR')
            self.assertIn('第 2 页', error.exception.error.message)
            self.assertEqual(len(self.events), 1)

    def test_docx_order_and_unsupported_images(self):
        self.require('python-docx')
        from docx import Document
        from docx.oxml import OxmlElement
        doc = Document()
        doc.add_paragraph('第一段')
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = '中间表格'
        doc.add_paragraph('末段知识点')
        path = self.root / 'ordered.docx'
        doc.save(path)
        text = reader.read_docx(path, self.context, time.monotonic() + 30)
        self.assertLess(text.index('第一段'), text.index('中间表格'))
        self.assertLess(text.index('中间表格'), text.index('末段知识点'))
        self.assertNotIn('第 1 页', text)
        doc.paragraphs[0]._p.append(OxmlElement('w:drawing'))
        doc.save(path)
        with self.assertRaises(PlatformError) as error:
            reader.read_docx(path, self.context, time.monotonic() + 30)
        self.assertIn('图片', error.exception.error.message)

    def test_docx_merged_and_nested_tables(self):
        self.require('python-docx')
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(1, 0)).text = '合并内容'
        cell = table.cell(0, 1)
        cell.text = '嵌套之前'
        cell.add_table(rows=1, cols=1).cell(0, 0).text = '嵌套内容'
        cell.add_paragraph('嵌套之后')
        path = self.root / 'nested.docx'
        doc.save(path)
        text = reader.read_docx(path, self.context, time.monotonic() + 30)
        self.assertEqual(text.count('合并内容'), 1)
        self.assertLess(text.index('嵌套之前'), text.index('嵌套内容'))
        self.assertLess(text.index('嵌套内容'), text.index('嵌套之后'))

    def test_text_pdf_blank_corrupt_encrypted_and_limits(self):
        self.require('PyMuPDF')
        import pymupdf
        path = self.root / 'text.pdf'
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((40, 50), 'FIRST PAGE')
            pdf.new_page()
            pdf.new_page().insert_text((40, 50), 'LAST PAGE UNIQUE FACT')
            pdf.save(path)
            pdf.save(self.root / 'encrypted.pdf', encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='secret', user_pw='password')
        with patch.object(reader, 'engine', side_effect=AssertionError('文字页无需 OCR')):
            text = reader.read_pdf(path, self.context, time.monotonic() + 30)
        self.assertIn('第 3 页', text)
        self.assertIn('LAST PAGE UNIQUE FACT', text)
        self.assertEqual(len(self.events), 3)
        with patch.object(reader, 'MAX_PAGES', 2), self.assertRaises(PlatformError):
            reader.read_pdf(path, self.context, time.monotonic() + 30)
        for name in ('corrupt.pdf', 'encrypted.pdf'):
            if name == 'corrupt.pdf': (self.root / name).write_bytes(b'broken')
            with self.assertRaises(PlatformError):
                reader.read_pdf(self.root / name, self.context, time.monotonic() + 30)
        with self.assertRaises(PlatformError):
            reader.read_pdf(path, self.context, time.monotonic() - 1)


if __name__ == '__main__':
    unittest.main()
