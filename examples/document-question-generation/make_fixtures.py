"""生成自制验收资料；须在已准备的文档块环境运行，不下载任何资源。"""
import argparse
from pathlib import Path
import pymupdf
from docx import Document

FACTS = [
    'Evaporation changes liquid water into water vapor.',
    'Condensation changes cooled water vapor into liquid water.',
    'Precipitation returns water from clouds to the ground.',
]


def make_fixtures(root):
    root.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph(FACTS[0])
    doc.add_table(rows=1, cols=1).cell(0, 0).text = FACTS[1]
    doc.add_paragraph(FACTS[2] + ' LAST PAGE UNIQUE FACT')
    doc.save(root / 'ordered.docx')
    with pymupdf.open() as text, pymupdf.open() as scan, pymupdf.open() as mixed:
        for index, fact in enumerate(FACTS):
            page = text.new_page(width=700, height=220)
            page.insert_text((25, 95), fact, fontsize=15)
            if index == 2:
                page.insert_text((25, 140), 'LAST PAGE UNIQUE FACT', fontsize=18)
            image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes('png')
            scan.new_page(width=700, height=220).insert_image(page.rect, stream=image)
            target = mixed.new_page(width=700, height=300)
            target.insert_text((25, 25), 'VISIBLE HEADER', fontsize=16)
            target.insert_image(pymupdf.Rect(0, 60, 700, 280), stream=image)
        text.save(root / 'text.pdf')
        scan.save(root / 'scan.pdf')
        mixed.save(root / 'mixed.pdf')
    with pymupdf.open() as blank:
        blank.new_page()
        blank.save(root / 'blank.pdf')
    (root / 'corrupt.pdf').write_bytes(b'%PDF-1.7\nbroken')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    make_fixtures(parser.parse_args().directory)
